# -*- coding: utf-8 -*-
"""
对汽车数据 Excel（.xlsx）全量清洗：新增一列【清洗后车型名称】。

清洗规则（按你给的拼接规则）：
年份款型 + 品牌车系 + 动力单元 + 变速箱/驱动形式 + 配置等级
示例：2019款 起亚KX5 1.6T 自动两驱豪华版

适配你这类“字段：值 + 车型全称 + 品牌路径”的表：
- 年款：2021
- 排量：2.9T
- 变速器：手自一体变速器(AT)
- 车型全称：2021款 阿尔法.罗密欧 朱丽叶GIULIA 2.9T 手自一体GTA
- 品牌路径（多行+>）：阿尔法.罗密欧\n>\n阿尔法.罗密欧_朱丽叶GIULIA\n>...

特点：
- 流式读取 openpyxl(read_only) + 流式写出 xlsxwriter(constant_memory)，可处理大表
- 默认处理所有工作表；每行末尾追加新列
- 默认认为第一行也是数据（无表头）。如果你的表第一行是表头，加 --has-header

依赖：
    openpyxl, xlsxwriter, tqdm
安装：
    conda install -y -c conda-forge openpyxl xlsxwriter tqdm
或：
    python -m pip install -U openpyxl xlsxwriter tqdm

用法：
    python clean_vehicle_name_v3.py --input 总3-1.xlsx --output 总3-1_清洗后.xlsx
抽样验收：
    python clean_vehicle_name_v3.py --input 总3-1.xlsx --output demo.xlsx --limit 10000
如果第一行是表头：
    python clean_vehicle_name_v3.py --input xxx.xlsx --output out.xlsx --has-header
"""

from __future__ import annotations
import argparse
import re
import sys
from pathlib import Path
from typing import Any, List, Optional

COLON = "："

# 依赖检查：缺啥就明确提示
try:
    import openpyxl  # type: ignore
except Exception:
    print("缺少依赖 openpyxl。执行：conda install -y -c conda-forge openpyxl  或  python -m pip install -U openpyxl",
          file=sys.stderr)
    raise

try:
    import xlsxwriter  # type: ignore
except Exception:
    print("缺少依赖 xlsxwriter。执行：conda install -y -c conda-forge xlsxwriter  或  python -m pip install -U xlsxwriter",
          file=sys.stderr)
    raise

try:
    from tqdm import tqdm  # type: ignore
except Exception:
    print("缺少依赖 tqdm。执行：conda install -y -c conda-forge tqdm  或  python -m pip install -U tqdm",
          file=sys.stderr)
    raise


def normalize_space(s: str) -> str:
    s = re.sub(r"[\u3000\t]+", " ", s)   # 全角空格/Tab
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def extract_full_name(cells: List[Any]) -> Optional[str]:
    """
    从一行里挑“车型全称”：
    - 含“款”
    - 不含“：”
    - 排除品牌路径（通常含换行+>）
    """
    best = None
    for c in cells:
        if not isinstance(c, str):
            continue
        if ("款" not in c) or (COLON in c):
            continue
        if "\n" in c and ">" in c:
            continue
        if best is None or len(c) > len(best):
            best = c
    return normalize_space(best) if best else None


def extract_year(year_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    year = None
    if year_raw:
        m = re.search(r"(\d{4})", str(year_raw))
        if m:
            year = m.group(1)
    if not year and full_name:
        m = re.match(r"\s*(\d{4})\s*款", full_name)
        if m:
            year = m.group(1)
    return f"{year}款" if year else None


def normalize_power(power_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    # 优先：排量/动力单元 字段
    if power_raw:
        p = str(power_raw)
        if COLON in p:
            p = p.split(COLON, 1)[1]
        p = normalize_space(p).replace(" ", "")
        p = re.sub(r"(?i)l\b", "L", p)
        p = re.sub(r"(?i)t\b", "T", p)
        return p or None

    # 兜底：从车型全称中抓 1.6T / 2.0L
    if full_name:
        m = re.search(r"(\d+(?:\.\d+)?\s*[TL])", full_name, flags=re.I)
        if m:
            return m.group(1).replace(" ", "").upper()
    return None


def normalize_transmission(trans_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    cand = None
    if trans_raw:
        t = str(trans_raw)
        if COLON in t:
            t = t.split(COLON, 1)[1]
        cand = normalize_space(t)
    elif full_name:
        cand = full_name

    if not cand:
        return None

    t = cand

    if re.search(r"手自一体", t):
        return "手自一体"

    m = re.search(r"(\d+)\s*AT", t, flags=re.I)
    if m:
        return f"{m.group(1)}AT"

    if re.search(r"CVT|无级", t, flags=re.I):
        return "CVT"
    if re.search(r"双离合|DCT", t, flags=re.I):
        return "双离合"
    if re.search(r"AMT", t, flags=re.I):
        return "AMT"
    if re.search(r"手动|MT", t, flags=re.I):
        return "手动"
    if re.search(r"自动|AT", t, flags=re.I):
        return "自动"

    t = re.sub(r"[()（）]", "", t)
    return normalize_space(t) or None


def normalize_drive(drive_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    s = ""
    if drive_raw:
        s = str(drive_raw)
        if COLON in s:
            s = s.split(COLON, 1)[1]
    elif full_name:
        s = full_name
    else:
        return None

    if re.search(r"全时四驱|适时四驱|分时四驱|四驱|4WD|AWD", s, flags=re.I):
        return "四驱"
    if re.search(r"两驱|2WD", s, flags=re.I):
        return "两驱"
    if re.search(r"前驱|后驱", s):
        return "两驱"
    return None


def parse_brand_series_from_path(path: Optional[str]) -> Optional[str]:
    """品牌路径（多行 + >）兜底抽品牌车系。"""
    if not path or not isinstance(path, str):
        return None

    parts = []
    for line in path.splitlines():
        line = line.strip()
        if not line or line == ">":
            continue
        if line.startswith(">"):
            continue
        parts.append(line)

    if not parts:
        return None

    brand = parts[0]
    series = None
    for p in parts[1:]:
        if "_" in p:
            series = p.split("_")[-1]
            break
    if not series and len(parts) >= 2:
        series = parts[1].split("_")[-1]

    if series:
        return f"{brand}{series}".replace(" ", "")
    return brand.replace(" ", "")


def extract_brand_series(full_name: Optional[str], power: Optional[str], brand_path: Optional[str]) -> Optional[str]:
    """
    优先从车型全称中截取（去掉年款后，到动力单元之前）。
    截不到再用品牌路径兜底。
    """
    if full_name:
        s = re.sub(r"^\s*\d{4}\s*款", "", full_name)
        s = normalize_space(s)

        if power:
            idx = s.find(power)
            if idx == -1:
                m = re.search(re.escape(power), s, flags=re.I)
                idx = m.start() if m else -1
            if idx > 0:
                pre = s[:idx].strip()
                return re.sub(r"\s+", "", pre)

    return parse_brand_series_from_path(brand_path)


def extract_trim(full_name: Optional[str],
                 power: Optional[str],
                 trans: Optional[str],
                 drive: Optional[str]) -> Optional[str]:
    """
    提取配置等级/版本：从车型全称剥离年款、动力、变速器、驱动后的剩余部分。
    """
    if not full_name:
        return None

    s = re.sub(r"^\s*\d{4}\s*款", "", full_name)
    s = normalize_space(s)

    # 找到动力单元token开始的位置，把前面当“品牌车系”范围不再细分
    tokens = s.split(" ")
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if power and (power.upper() in tok.replace(" ", "").upper()):
            break
        if re.fullmatch(r"\d+(?:\.\d+)?[TL]", tok, flags=re.I):
            break
        i += 1
    rest = " ".join(tokens[i:])

    if power:
        rest = re.sub(re.escape(power), " ", rest, flags=re.I)

    # 噪声：马力/功率/扭矩
    rest = re.sub(r"\d+\s*(PS|马力|kW|KW|Nm|N·m)", " ", rest, flags=re.I)

    # 噪声：发动机技术标识（独立token）
    rest = re.sub(r"(?<![A-Za-z0-9])(SC|TSI|TFSI|GDI|CVVT|DVVT|VVT|VTEC|i-VTEC|EcoBoost)(?![A-Za-z0-9])",
                  " ", rest, flags=re.I)

    if trans:
        rest = re.sub(re.escape(trans), " ", rest, flags=re.I)
    if drive:
        rest = re.sub(re.escape(drive), " ", rest, flags=re.I)

    rest = re.sub(r"两驱|四驱|前驱|后驱|AWD|4WD|2WD", " ", rest, flags=re.I)
    rest = re.sub(r"\d+\s*(挡|速)", " ", rest)

    out = normalize_space(rest)
    return out or None


def build_clean_name(cells: List[Any]) -> Optional[str]:
    year_raw = None
    power_raw = None
    trans_raw = None
    drive_raw = None
    brand_path = None

    for c in cells:
        if not isinstance(c, str):
            continue
        if c.startswith("年款：") or c.startswith("年款:"):
            year_raw = c
        elif c.startswith("排量：") or c.startswith("排量:") or c.startswith("动力单元：") or c.startswith("动力单元:"):
            power_raw = c
        elif c.startswith("变速器：") or c.startswith("变速器:") or c.startswith("变速箱：") or c.startswith("变速箱:"):
            trans_raw = c
        elif c.startswith("驱动：") or c.startswith("驱动:") or c.startswith("驱动形式：") or c.startswith("驱动形式:"):
            drive_raw = c
        elif "\n" in c and ">" in c:
            brand_path = c

    full_name = extract_full_name(cells)
    year_part = extract_year(year_raw, full_name)
    power = normalize_power(power_raw, full_name)
    trans = normalize_transmission(trans_raw, full_name)
    drive = normalize_drive(drive_raw, full_name)
    brand_series = extract_brand_series(full_name, power, brand_path)
    trim = extract_trim(full_name, power, trans, drive)

    # 末段：变速箱/驱动形式 + 配置等级（不加空格，和示例一致）
    last = ""
    if trans:
        last += trans
    if drive:
        last += drive
    if trim:
        last += trim

    parts = []
    if year_part:
        parts.append(year_part)
    if brand_series:
        parts.append(brand_series)
    if power:
        parts.append(power)
    if last:
        parts.append(last)

    return " ".join(parts) if parts else None


def process_file(input_path: Path, output_path: Path, limit: int = 0, has_header: bool = False) -> None:
    wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)

    out = xlsxwriter.Workbook(str(output_path), {
        "constant_memory": True,
        "strings_to_urls": False,
        "nan_inf_to_errors": True,
    })

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        wso = out.add_worksheet(sheet_name[:31])

        total = ws.max_row if ws.max_row and ws.max_row > 0 else None
        pbar = tqdm(total=min(total, limit) if (total and limit) else (total or None),
                    desc=f"处理 {sheet_name}", unit="行")

        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if limit and r_idx >= limit:
                break

            cells = list(row)

            # 表头处理：如果第一行是表头，则只追加标题，不解析
            if has_header and r_idx == 0:
                cells.append("清洗后车型名称")
            else:
                cells.append(build_clean_name(cells))

            wso.write_row(r_idx, 0, cells)
            pbar.update(1)

        pbar.close()

    out.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="输入Excel路径（.xlsx）")
    ap.add_argument("--output", required=True, help="输出Excel路径（.xlsx）")
    ap.add_argument("--limit", type=int, default=0, help="只处理前N行用于验收（0=全量）")
    ap.add_argument("--has-header", action="store_true", help="第一行是表头时启用（会在第一行写入新列标题）")
    args = ap.parse_args()

    process_file(Path(args.input), Path(args.output), limit=args.limit, has_header=args.has_header)


if __name__ == "__main__":
    main()
