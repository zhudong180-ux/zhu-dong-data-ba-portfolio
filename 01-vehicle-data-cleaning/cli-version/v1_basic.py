# -*- coding: utf-8 -*-
"""
为超大Excel（百万行级）新增一列“清洗后车型名称”。

清洗规则（你给的规则）：
年份款型 + 品牌车系 + 动力单元 + 变速箱/驱动形式 + 配置等级
示例：2019款 起亚KX5 1.6T 自动两驱豪华版

本脚本特点：
- 流式读取（openpyxl read_only）+ 流式写出（xlsxwriter constant_memory），适合百万行
- 默认“保留原表所有单元格原样”，仅在末尾新增一列
- 可选：把类似“排量：3.0T”这种单元格去掉前缀，只保留值（--strip-prefix）

用法示例：
1）只新增一列（不改动原单元格）
    python clean_vehicle_name.py --input 总1-1.xlsx --output 总1-1_清洗后.xlsx

2）新增一列 + 顺便把所有“字段：值”改成“值”
    python clean_vehicle_name.py --input 总1-1.xlsx --output 总1-1_清洗后.xlsx --strip-prefix

3）只跑前10000行做验证
    python clean_vehicle_name.py --input 总1-1.xlsx --output demo.xlsx --limit 10000
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any, Iterable, List, Optional, Tuple

import openpyxl
import xlsxwriter
from tqdm import tqdm


COLON = "："


def normalize_space(s: str) -> str:
    s = re.sub(r"[\u3000\t]+", " ", s)   # 全角空格/Tab
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def strip_prefix_if_kv(cell: Any) -> Any:
    """把 '字段：值' 变成 '值'（可选清洗项）。"""
    if isinstance(cell, str) and COLON in cell:
        _, v = cell.split(COLON, 1)
        return v.strip()
    return cell


def extract_full_name(cells: List[Any]) -> Optional[str]:
    """
    在一行里选出最像“车型全称”的单元格：
    - 包含“款”
    - 不包含“：”
    - 排除品牌路径那种带换行和“>”的单元格
    """
    best = None
    for c in cells:
        if isinstance(c, str) and ("款" in c) and (COLON not in c):
            if "\n" in c and ">" in c:
                continue
            if best is None or len(c) > len(best):
                best = c
    return normalize_space(best) if best else None


def extract_year_from_any(year_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    year = None
    if year_raw:
        m = re.search(r"(\d{4})", str(year_raw))
        if m:
            year = m.group(1)
    if not year and full_name:
        m = re.match(r"\s*(\d{4})\s*款", full_name)
        if m:
            year = m.group(1)
    return f"{year}款" if year else None


def normalize_power(power_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    # 优先使用“排量：”字段
    if power_raw:
        p = str(power_raw)
        if COLON in p:
            p = p.split(COLON, 1)[1]
        p = normalize_space(p).replace(" ", "")
        p = re.sub(r"(?i)l\b", "L", p)
        p = re.sub(r"(?i)t\b", "T", p)
        return p or None

    # 兜底：从车型全称里抓一个 1.6T / 2.0L 之类
    if full_name:
        m = re.search(r"(\d+(?:\.\d+)?\s*[TL])", full_name, flags=re.I)
        if m:
            p = m.group(1).replace(" ", "").upper()
            return p
    return None


def normalize_transmission(trans_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    cand = None
    if trans_raw:
        t = str(trans_raw)
        if COLON in t:
            t = t.split(COLON, 1)[1]
        cand = normalize_space(t)
    elif full_name:
        cand = full_name

    if not cand:
        return None

    t = cand
    if re.search(r"手自一体", t):
        return "手自一体"

    m = re.search(r"(\d+)\s*AT", t, flags=re.I)
    if m:
        return f"{m.group(1)}AT"

    if re.search(r"CVT|无级", t, flags=re.I):
        return "CVT"

    if re.search(r"双离合|DCT", t, flags=re.I):
        return "双离合"

    if re.search(r"AMT", t, flags=re.I):
        return "AMT"

    if re.search(r"手动|MT", t, flags=re.I):
        return "手动"

    if re.search(r"自动|AT", t, flags=re.I):
        return "自动"

    # 兜底：去掉括号等
    t = re.sub(r"[()（）]", "", t)
    t = normalize_space(t)
    return t or None


def normalize_drive(drive_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    s = ""
    if drive_raw:
        s = str(drive_raw)
        if COLON in s:
            s = s.split(COLON, 1)[1]
    elif full_name:
        s = full_name
    else:
        return None

    if re.search(r"全时四驱|适时四驱|分时四驱|四驱|4WD|AWD", s, flags=re.I):
        return "四驱"
    if re.search(r"两驱|2WD", s, flags=re.I):
        return "两驱"
    if re.search(r"前驱|后驱", s):
        return "两驱"
    return None


def parse_brand_series_from_path(path: Optional[str]) -> Optional[str]:
    if not path or not isinstance(path, str):
        return None

    parts = []
    for line in path.splitlines():
        line = line.strip()
        if not line or line == ">":
            continue
        if line.startswith(">"):
            continue
        parts.append(line)

    if not parts:
        return None

    brand = parts[0]
    series = None
    for p in parts[1:]:
        if "_" in p:
            series = p.split("_")[-1]
            break
    if not series and len(parts) >= 2:
        series = parts[1].split("_")[-1]

    if series:
        return f"{brand}{series}".replace(" ", "")
    return brand.replace(" ", "")


def extract_brand_series(full_name: Optional[str], power: Optional[str], brand_path: Optional[str]) -> Optional[str]:
    if full_name:
        s = re.sub(r"^\s*\d{4}\s*款", "", full_name)
        s = normalize_space(s)

        if power:
            idx = s.find(power)
            if idx == -1:
                m = re.search(re.escape(power), s, flags=re.I)
                idx = m.start() if m else -1
            if idx > 0:
                pre = s[:idx].strip()
                return re.sub(r"\s+", "", pre)

    return parse_brand_series_from_path(brand_path)


def extract_trim(full_name: Optional[str],
                 power: Optional[str],
                 trans: Optional[str],
                 drive: Optional[str]) -> Optional[str]:
    if not full_name:
        return None

    s = re.sub(r"^\s*\d{4}\s*款", "", full_name)
    s = normalize_space(s)

    # 移除“品牌车系”（假设在最前面直到遇到动力单元）
    tokens = s.split(" ")
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if power and (power.upper() in tok.replace(" ", "").upper()):
            break
        if re.fullmatch(r"\d+(?:\.\d+)?[TL]", tok, flags=re.I):
            break
        i += 1
    rest = " ".join(tokens[i:])

    if power:
        rest = re.sub(re.escape(power), " ", rest, flags=re.I)

    # 去掉马力/功率/扭矩之类
    rest = re.sub(r"\d+\s*(PS|马力|kW|KW|Nm|N·m)", " ", rest, flags=re.I)

    # 去掉常见发动机技术码（只在“独立token”时剔除，避免误删）
    rest = re.sub(r"(?<![A-Za-z0-9])(SC|TSI|TFSI|GDI|CVVT|DVVT|VVT|VTEC|i-VTEC|EcoBoost)(?![A-Za-z0-9])",
                  " ", rest, flags=re.I)

    if trans:
        rest = re.sub(re.escape(trans), " ", rest, flags=re.I)

    if drive:
        rest = re.sub(re.escape(drive), " ", rest, flags=re.I)

    rest = re.sub(r"两驱|四驱|前驱|后驱|AWD|4WD|2WD", " ", rest, flags=re.I)
    rest = re.sub(r"\d+\s*(挡|速)", " ", rest)

    return normalize_space(rest) or None


def build_clean_name(cells: List[Any]) -> Optional[str]:
    year_raw = None
    power_raw = None
    trans_raw = None
    drive_raw = None
    brand_path = None

    for c in cells:
        if isinstance(c, str):
            if c.startswith("年款：") or c.startswith("年款:"):
                year_raw = c
            elif c.startswith("排量：") or c.startswith("排量:") or c.startswith("动力单元：") or c.startswith("动力单元:"):
                power_raw = c
            elif c.startswith("变速器：") or c.startswith("变速器:") or c.startswith("变速箱：") or c.startswith("变速箱:"):
                trans_raw = c
            elif c.startswith("驱动：") or c.startswith("驱动:") or c.startswith("驱动形式：") or c.startswith("驱动形式:"):
                drive_raw = c
            elif "\n" in c and ">" in c:
                brand_path = c

    full_name = extract_full_name(cells)
    year_part = extract_year_from_any(year_raw, full_name)
    power = normalize_power(power_raw, full_name)
    trans = normalize_transmission(trans_raw, full_name)
    drive = normalize_drive(drive_raw, full_name)
    brand_series = extract_brand_series(full_name, power, brand_path)
    trim = extract_trim(full_name, power, trans, drive)

    # 按规则拼接：最后一段是“变速箱/驱动形式 + 配置等级”（通常不加空格）
    last = ""
    if trans:
        last += trans
    if drive:
        last += drive
    if trim:
        last += trim

    parts = []
    if year_part:
        parts.append(year_part)
    if brand_series:
        parts.append(brand_series)
    if power:
        parts.append(power)
    if last:
        parts.append(last)

    return " ".join(parts) if parts else None


def process_file(input_path: Path,
                 output_path: Path,
                 strip_prefix: bool = False,
                 limit: int = 0,
                 sheets: Optional[List[str]] = None) -> None:
    wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)

    # 写出使用 xlsxwriter，速度更快、内存更稳
    out = xlsxwriter.Workbook(output_path, {
        "constant_memory": True,
        "strings_to_urls": False,
        "nan_inf_to_errors": True,
    })

    target_sheets = sheets or wb.sheetnames
    for sheet_name in target_sheets:
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        wso = out.add_worksheet(sheet_name[:31])  # Excel sheet name 限制 31

        total = ws.max_row if ws.max_row and ws.max_row > 0 else None
        it = ws.iter_rows(values_only=True)

        row_idx = 0
        pbar = tqdm(total=min(total, limit) if (total and limit) else (total or None),
                    desc=f"处理工作表 {sheet_name}", unit="行")

        for row in it:
            if limit and row_idx >= limit:
                break

            cells = list(row)
            clean_name = build_clean_name(cells)

            if strip_prefix:
                cells = [strip_prefix_if_kv(c) for c in cells]

            # 追加一列
            cells.append(clean_name)

            # 写一行
            wso.write_row(row_idx, 0, cells)
            row_idx += 1
            pbar.update(1)

        pbar.close()

    out.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="输入Excel路径（.xlsx）")
    ap.add_argument("--output", required=True, help="输出Excel路径（.xlsx）")
    ap.add_argument("--strip-prefix", action="store_true",
                    help="把类似 '排量：3.0T' 的单元格改成 '3.0T'（可选）")
    ap.add_argument("--limit", type=int, default=0,
                    help="只处理前N行（0表示全量处理），用于快速验证")
    ap.add_argument("--sheets", nargs="*", default=None,
                    help="只处理指定工作表（不填则处理全部）")
    args = ap.parse_args()

    process_file(
        input_path=Path(args.input),
        output_path=Path(args.output),
        strip_prefix=args.strip_prefix,
        limit=args.limit,
        sheets=args.sheets
    )


if __name__ == "__main__":
    main()
