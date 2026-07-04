# -*- coding: utf-8 -*-
"""
用途：对汽车数据 Excel（.xlsx）全量清洗：逐行生成“清洗后车型名称”，并在末尾新增一列写入到新文件。

车型名称规则：
年份款型 + 品牌车系 + 动力单元 + 变速箱/驱动形式 + 配置等级
示例：2019款 起亚KX5 1.6T 自动两驱豪华版

本脚本适配两类常见源数据：
A) “字段：值”散列在多个单元格（年款：/排量：/变速器：/驱动：） + 单独的“车型全称”单元格
B) “品牌路径”单元格（多行 + '>'）里最后一行就是车型全称（你这份“总表_精简版”就是这种）

特点：
- openpyxl(read_only) 流式读 + xlsxwriter(constant_memory) 流式写，适合大表
- 默认第一行也是数据（无表头）；如第一行是表头，用 --has-header
- 默认处理所有工作表；可用 --sheets 指定

依赖：
    openpyxl, xlsxwriter, tqdm
安装：
    conda install -y -c conda-forge openpyxl xlsxwriter tqdm
或：
    python -m pip install -U openpyxl xlsxwriter tqdm

用法：
    python clean_vehicle_name_v4.py --input "总表_ 精简版 -1.xlsx" --output "总表_精简版_清洗后.xlsx"
抽样验收：
    python clean_vehicle_name_v4.py --input in.xlsx --output demo.xlsx --limit 10000
"""

from __future__ import annotations
import argparse
import re
import sys
from pathlib import Path
from typing import Any, List, Optional

COLON = "："

# 依赖检查：缺啥就明确提示（不让你猜）
try:
    import openpyxl  # type: ignore
except Exception:
    print("缺少依赖 openpyxl。执行：conda install -y -c conda-forge openpyxl  或  python -m pip install -U openpyxl",
          file=sys.stderr)
    raise

try:
    import xlsxwriter  # type: ignore
except Exception:
    print("缺少依赖 xlsxwriter。执行：conda install -y -c conda-forge xlsxwriter  或  python -m pip install -U xlsxwriter",
          file=sys.stderr)
    raise

try:
    from tqdm import tqdm  # type: ignore
except Exception:
    print("缺少依赖 tqdm。执行：conda install -y -c conda-forge tqdm  或  python -m pip install -U tqdm",
          file=sys.stderr)
    raise


def normalize_space(s: str) -> str:
    s = re.sub(r"[\u3000\t]+", " ", s)   # 全角空格/Tab
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def get_brand_path_cell(cells: List[Any]) -> Optional[str]:
    """品牌路径单元格常见特征：含换行且含 '>'"""
    for c in cells:
        if isinstance(c, str) and ("\n" in c) and (">" in c):
            return c
    return None


def extract_full_name(cells: List[Any]) -> Optional[str]:
    """
    优先：从独立单元格中找车型全称（含“款”，不含“：”，且不是品牌路径）
    兜底：若找不到，则从品牌路径单元格最后一段抽取车型全称
    """
    best = None
    for c in cells:
        if not isinstance(c, str):
            continue
        if ("款" not in c) or (COLON in c):
            continue
        if "\n" in c and ">" in c:
            continue
        if best is None or len(c) > len(best):
            best = c
    if best:
        return normalize_space(best)

    # 兜底：品牌路径最后一行通常是车型全称（例如：>2001款铃木奥拓0.8L手动基本型）
    bp = get_brand_path_cell(cells)
    if not bp:
        return None

    lines = [ln.strip() for ln in bp.splitlines() if ln.strip()]
    # 过滤掉纯 ">"
    lines = [ln for ln in lines if ln != ">"]
    # 从后往前找包含“款”的行
    for ln in reversed(lines):
        # 可能带前导 '>'，去掉
        ln2 = ln.lstrip(">").strip()
        if "款" in ln2 and COLON not in ln2:
            return normalize_space(ln2)
    return None


def extract_year(year_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    year = None
    if year_raw:
        m = re.search(r"(\d{4})", str(year_raw))
        if m:
            year = m.group(1)
    if not year and full_name:
        m = re.match(r"\s*(\d{4})\s*款", full_name)
        if m:
            year = m.group(1)
    return f"{year}款" if year else None


def normalize_power(power_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    # 优先：排量/动力单元 字段
    if power_raw:
        p = str(power_raw)
        if COLON in p:
            p = p.split(COLON, 1)[1]
        p = normalize_space(p).replace(" ", "")
        p = re.sub(r"(?i)l\b", "L", p)
        p = re.sub(r"(?i)t\b", "T", p)
        return p or None

    # 兜底：从车型全称中抓 1.6T / 2.0L
    if full_name:
        m = re.search(r"(\d+(?:\.\d+)?\s*[TL])", full_name, flags=re.I)
        if m:
            return m.group(1).replace(" ", "").upper()
    return None


def normalize_transmission(trans_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    cand = None
    if trans_raw:
        t = str(trans_raw)
        if COLON in t:
            t = t.split(COLON, 1)[1]
        cand = normalize_space(t)
    elif full_name:
        cand = full_name

    if not cand:
        return None

    t = cand

    # 明确类别优先
    if re.search(r"手自一体", t):
        return "手自一体"

    m = re.search(r"(\d+)\s*AT", t, flags=re.I)
    if m:
        return f"{m.group(1)}AT"

    if re.search(r"CVT|无级", t, flags=re.I):
        return "CVT"
    if re.search(r"双离合|DCT", t, flags=re.I):
        return "双离合"
    if re.search(r"AMT", t, flags=re.I):
        return "AMT"
    if re.search(r"手动|MT", t, flags=re.I):
        return "手动"
    if re.search(r"自动|AT", t, flags=re.I):
        return "自动"

    t = re.sub(r"[()（）]", "", t)
    return normalize_space(t) or None


def normalize_drive(drive_raw: Optional[str], full_name: Optional[str]) -> Optional[str]:
    s = ""
    if drive_raw:
        s = str(drive_raw)
        if COLON in s:
            s = s.split(COLON, 1)[1]
    elif full_name:
        s = full_name
    else:
        return None

    if re.search(r"全时四驱|适时四驱|分时四驱|四驱|4WD|AWD", s, flags=re.I):
        return "四驱"
    if re.search(r"两驱|2WD", s, flags=re.I):
        return "两驱"
    if re.search(r"前驱|后驱", s):
        return "两驱"
    return None


def parse_brand_series_from_path(path: Optional[str]) -> Optional[str]:
    """品牌路径（多行 + >）兜底抽品牌车系。"""
    if not path or not isinstance(path, str):
        return None

    parts = []
    for line in path.splitlines():
        line = line.strip()
        if not line or line == ">":
            continue
        if line.startswith(">"):
            continue
        parts.append(line)

    if not parts:
        return None

    brand = parts[0]
    series = None
    for p in parts[1:]:
        if "_" in p:
            series = p.split("_")[-1]
            break
    if not series and len(parts) >= 2:
        series = parts[1].split("_")[-1]

    if series:
        return f"{brand}{series}".replace(" ", "")
    return brand.replace(" ", "")


def extract_brand_series(full_name: Optional[str], power: Optional[str], brand_path: Optional[str]) -> Optional[str]:
    """
    优先从车型全称中截取（去掉年款后，到动力单元之前）。
    注意：车型全称可能无空格（例如：2002款比亚迪福莱尔0.8L手动基本型），此时仍可用“动力单元”定位截取。
    """
    if full_name:
        s = re.sub(r"^\s*\d{4}\s*款", "", full_name)
        s = normalize_space(s)

        if power:
            idx = s.find(power)
            if idx == -1:
                m = re.search(re.escape(power), s, flags=re.I)
                idx = m.start() if m else -1
            if idx > 0:
                pre = s[:idx].strip()
                # 去掉空格
                return re.sub(r"\s+", "", pre)

    return parse_brand_series_from_path(brand_path)


def extract_trim(full_name: Optional[str],
                 brand_series: Optional[str],
                 power: Optional[str],
                 trans: Optional[str],
                 drive: Optional[str]) -> Optional[str]:
    """
    从车型全称提取“配置等级/版本”。
    关键：兼容“无空格”的车型全称（直接做字符串剥离）。
    """
    if not full_name:
        return None

    s = full_name
    s = re.sub(r"^\s*\d{4}\s*款", "", s)
    s = normalize_space(s).replace(" ", "")  # 强制无空格统一处理

    if brand_series:
        bs = str(brand_series).replace(" ", "")
        if s.startswith(bs):
            s = s[len(bs):]
        else:
            # 有些全称包含“厂商+车系”，品牌路径可能更准，尽量不硬删
            pass

    if power:
        s = re.sub(re.escape(power), "", s, flags=re.I)

    # 去掉变速箱关键词（用更广的模式）
    # trans 可能是“手动/自动/手自一体/8AT...”
    if trans:
        s = re.sub(re.escape(trans), "", s, flags=re.I)

    # 去掉驱动
    if drive:
        s = re.sub(re.escape(drive), "", s)

    # 去掉常见噪声：马力/功率/扭矩/发动机技术标识
    s = re.sub(r"\d+(PS|马力|kW|KW|Nm|N·m)", "", s)
    s = re.sub(r"(SC|TSI|TFSI|GDI|CVVT|DVVT|VVT|VTEC|i-VTEC|EcoBoost)", "", s, flags=re.I)

    # 清理剩余的“两驱/四驱/前驱/后驱”等（以免 drive 字段缺失时残留）
    s = re.sub(r"(两驱|四驱|前驱|后驱|AWD|4WD|2WD)", "", s, flags=re.I)

    # 清理挡位类
    s = re.sub(r"\d+(挡|速)", "", s)

    out = normalize_space(s)
    return out or None


def build_clean_name(cells: List[Any]) -> Optional[str]:
    year_raw = None
    power_raw = None
    trans_raw = None
    drive_raw = None

    for c in cells:
        if not isinstance(c, str):
            continue
        if c.startswith("年款：") or c.startswith("年款:"):
            year_raw = c
        elif c.startswith("排量：") or c.startswith("排量:") or c.startswith("动力单元：") or c.startswith("动力单元:"):
            power_raw = c
        elif c.startswith("变速器：") or c.startswith("变速器:") or c.startswith("变速箱：") or c.startswith("变速箱:"):
            trans_raw = c
        elif c.startswith("驱动：") or c.startswith("驱动:") or c.startswith("驱动形式：") or c.startswith("驱动形式:"):
            drive_raw = c

    brand_path = get_brand_path_cell(cells)
    full_name = extract_full_name(cells)

    year_part = extract_year(year_raw, full_name)
    power = normalize_power(power_raw, full_name)
    trans = normalize_transmission(trans_raw, full_name)
    drive = normalize_drive(drive_raw, full_name)
    brand_series = extract_brand_series(full_name, power, brand_path)
    trim = extract_trim(full_name, brand_series, power, trans, drive)

    # 末段：变速箱/驱动形式 + 配置等级（不加空格，符合示例）
    last = ""
    if trans:
        last += trans
    if drive:
        last += drive
    if trim:
        last += trim

    parts = []
    if year_part:
        parts.append(year_part)
    if brand_series:
        parts.append(brand_series)
    if power:
        parts.append(power)
    if last:
        parts.append(last)

    return " ".join(parts) if parts else None


def process_file(input_path: Path,
                 output_path: Path,
                 limit: int = 0,
                 has_header: bool = False,
                 sheets: Optional[List[str]] = None,
                 new_col_name: str = "清洗后车型名称") -> None:
    wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)

    out = xlsxwriter.Workbook(str(output_path), {
        "constant_memory": True,
        "strings_to_urls": False,
        "nan_inf_to_errors": True,
    })

    target_sheets = sheets or wb.sheetnames
    for sheet_name in target_sheets:
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        wso = out.add_worksheet(sheet_name[:31])

        total = ws.max_row if ws.max_row and ws.max_row > 0 else None
        pbar = tqdm(total=min(total, limit) if (total and limit) else (total or None),
                    desc=f"处理 {sheet_name}", unit="行")

        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if limit and r_idx >= limit:
                break

            cells = list(row)
            if has_header and r_idx == 0:
                cells.append(new_col_name)
            else:
                cells.append(build_clean_name(cells))

            wso.write_row(r_idx, 0, cells)
            pbar.update(1)

        pbar.close()

    out.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="输入Excel路径（.xlsx）")
    ap.add_argument("--output", required=True, help="输出Excel路径（.xlsx）")
    ap.add_argument("--limit", type=int, default=0, help="只处理前N行用于验收（0=全量）")
    ap.add_argument("--has-header", action="store_true", help="第一行是表头时启用（会在第一行写入新列标题）")
    ap.add_argument("--sheets", nargs="*", default=None, help="只处理指定工作表（不填=全部）")
    ap.add_argument("--new-col-name", default="清洗后车型名称", help="新增列的列名（仅在 --has-header 时写入第一行）")
    args = ap.parse_args()

    process_file(
        input_path=Path(args.input),
        output_path=Path(args.output),
        limit=args.limit,
        has_header=args.has_header,
        sheets=args.sheets,
        new_col_name=args.new_col_name
    )


if __name__ == "__main__":
    main()
